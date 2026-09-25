from pathlib import Path
import json

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# ============================================================
# PATHS
# ============================================================

BASE_DIR = Path(__file__).resolve().parent
DOCS_DIR = BASE_DIR / "DOCS"

BRAND_BOOK_PATH = DOCS_DIR / "brand-book.json"
OUTPUT_PATH = DOCS_DIR / "template-proposta.docx"


# ============================================================
# UTILITIES
# ============================================================

def hex_to_rgb(hex_color):
    """Converte #RRGGBB para RGBColor."""
    value = hex_color.lstrip("#")
    return RGBColor(
        int(value[0:2], 16),
        int(value[2:4], 16),
        int(value[4:6], 16),
    )


def set_cell_background(cell, fill):
    """Define a cor de fundo de uma célula."""
    tc_pr = cell._tc.get_or_add_tcPr()

    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)

    shd.set(qn("w:fill"), fill.replace("#", ""))


def set_cell_border(cell, color="E2E8F0", size="4"):
    """Aplica bordas discretas à célula."""
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()

    borders = tc_pr.first_child_found_in("w:tcBorders")

    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))

        if element is None:
            element = OxmlElement(tag)
            borders.append(element)

        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color.replace("#", ""))


def set_repeat_table_header(row):
    """Marca a primeira linha da tabela para repetição em novas páginas."""
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_bottom_border(paragraph, color="E2E8F0", size="6"):
    """Adiciona uma linha inferior ao parágrafo."""
    p = paragraph._p
    p_pr = p.get_or_add_pPr()

    p_bdr = p_pr.find(qn("w:pBdr"))

    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)

    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color.replace("#", ""))

    p_bdr.append(bottom)


def add_page_number(paragraph):
    """Insere campo PAGE no rodapé."""
    run = paragraph.add_run()

    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "

    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


# ============================================================
# LOAD BRAND BOOK
# ============================================================

if not BRAND_BOOK_PATH.exists():
    raise FileNotFoundError(
        f"Brand book não encontrado: {BRAND_BOOK_PATH}"
    )

with BRAND_BOOK_PATH.open("r", encoding="utf-8") as file:
    brand = json.load(file)


colors = brand["colors"]
typography = brand["typography"]
document_config = brand["document"]
sections_config = brand["sections"]["required"]


PRIMARY = colors["primary"]["hex"]
SECONDARY = colors["secondary"]["hex"]
ACCENT = colors["accent"]["hex"]
TEXT = colors["text"]["hex"]
MUTED = colors["muted_text"]["hex"]
LIGHT_BG = colors["light_background"]["hex"]
BORDER = colors["border"]["hex"]
WHITE = colors["white"]["hex"]

FONT = typography["primary_font"]


# ============================================================
# DOCUMENT
# ============================================================

doc = Document()

section = doc.sections[0]

section.top_margin = Cm(document_config["margins_cm"]["top"])
section.bottom_margin = Cm(document_config["margins_cm"]["bottom"])
section.left_margin = Cm(document_config["margins_cm"]["left"])
section.right_margin = Cm(document_config["margins_cm"]["right"])


# ============================================================
# GLOBAL STYLES
# ============================================================

normal_style = doc.styles["Normal"]
normal_style.font.name = FONT
normal_style.font.size = Pt(typography["body"]["font_size_pt"])
normal_style.font.color.rgb = hex_to_rgb(TEXT)

normal_style.paragraph_format.space_after = Pt(6)
normal_style.paragraph_format.line_spacing = typography["body"]["line_spacing"]


for style_name in ["Title", "Heading 1", "Heading 2"]:
    style = doc.styles[style_name]
    style.font.name = FONT


title_style = doc.styles["Title"]
title_style.font.size = Pt(typography["title"]["font_size_pt"])
title_style.font.bold = True
title_style.font.color.rgb = hex_to_rgb(PRIMARY)

heading1 = doc.styles["Heading 1"]
heading1.font.size = Pt(typography["section_heading"]["font_size_pt"])
heading1.font.bold = True
heading1.font.color.rgb = hex_to_rgb(PRIMARY)

heading2 = doc.styles["Heading 2"]
heading2.font.size = Pt(typography["subsection_heading"]["font_size_pt"])
heading2.font.bold = True
heading2.font.color.rgb = hex_to_rgb(SECONDARY)


# ============================================================
# HEADER
# ============================================================

header = section.header
header_p = header.paragraphs[0]

header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT

run = header_p.add_run("NEXORA CONSULTING")
run.font.name = FONT
run.font.size = Pt(8)
run.font.bold = True
run.font.color.rgb = hex_to_rgb(MUTED)

run = header_p.add_run("                                         DealCraft AI")
run.font.name = FONT
run.font.size = Pt(8)
run.font.color.rgb = hex_to_rgb(MUTED)

add_bottom_border(header_p, BORDER, "4")


# ============================================================
# FOOTER
# ============================================================

footer = section.footer
footer_p = footer.paragraphs[0]

footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = footer_p.add_run(
    "NEXORA CONSULTING · Data · AI · Growth · Automation   |   "
    "CONFIDENCIAL   |   Página "
)

run.font.name = FONT
run.font.size = Pt(8)
run.font.color.rgb = hex_to_rgb(MUTED)

add_page_number(footer_p)


# ============================================================
# COVER
# ============================================================

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(3)

run = p.add_run("NEXORA CONSULTING")
run.font.name = FONT
run.font.size = Pt(14)
run.font.bold = True
run.font.color.rgb = hex_to_rgb(PRIMARY)

p = doc.add_paragraph()

run = p.add_run("Data · AI · Growth · Automation")
run.font.name = FONT
run.font.size = Pt(9)
run.font.color.rgb = hex_to_rgb(MUTED)

doc.add_paragraph()
doc.add_paragraph()
doc.add_paragraph()

p = doc.add_paragraph()

run = p.add_run("PROPOSTA COMERCIAL")
run.font.name = FONT
run.font.size = Pt(10)
run.font.bold = True
run.font.color.rgb = hex_to_rgb(SECONDARY)

p = doc.add_paragraph()

run = p.add_run("[NOME DO CLIENTE]")
run.font.name = FONT
run.font.size = Pt(28)
run.font.bold = True
run.font.color.rgb = hex_to_rgb(PRIMARY)

p = doc.add_paragraph()

run = p.add_run("[TÍTULO / SOLUÇÃO PRINCIPAL]")
run.font.name = FONT
run.font.size = Pt(16)
run.font.bold = True
run.font.color.rgb = hex_to_rgb(TEXT)

doc.add_paragraph()
doc.add_paragraph()

metadata = [
    ("PROPOSTA", "[NUMERO]/[ANO]"),
    ("EMISSÃO", "[DD/MM/AAAA]"),
    ("VALIDADE", "30 dias"),
    ("CONSULTOR", "[NOME DO CONSULTOR]"),
]

table = doc.add_table(rows=len(metadata), cols=2)
table.alignment = WD_TABLE_ALIGNMENT.LEFT
table.autofit = True

for i, (label, value) in enumerate(metadata):
    left = table.cell(i, 0)
    right = table.cell(i, 1)

    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    set_cell_background(left, LIGHT_BG)

    for cell in (left, right):
        set_cell_border(cell, BORDER)

    p_left = left.paragraphs[0]
    r_left = p_left.add_run(label)
    r_left.font.name = FONT
    r_left.font.size = Pt(8)
    r_left.font.bold = True
    r_left.font.color.rgb = hex_to_rgb(MUTED)

    p_right = right.paragraphs[0]
    r_right = p_right.add_run(value)
    r_right.font.name = FONT
    r_right.font.size = Pt(9)
    r_right.font.color.rgb = hex_to_rgb(TEXT)


doc.add_page_break()


# ============================================================
# SECTION HELPERS
# ============================================================

def add_section_header(number, title):
    p = doc.add_paragraph()

    r = p.add_run(f"{number}  ")
    r.font.name = FONT
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = hex_to_rgb(ACCENT)

    r = p.add_run(title.upper())
    r.font.name = FONT
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = hex_to_rgb(PRIMARY)

    add_bottom_border(p, BORDER, "6")

    p.paragraph_format.space_after = Pt(14)


def add_placeholder(text):
    p = doc.add_paragraph()

    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(10.5)
    r.font.color.rgb = hex_to_rgb(MUTED)
    r.font.italic = True

    p.paragraph_format.space_after = Pt(8)


# ============================================================
# 01 — CONTEXTO E OPORTUNIDADE
# ============================================================

add_section_header("01", "Contexto e oportunidade")

add_placeholder(
    "[Descrever o cenário atual do cliente, contexto do negócio, "
    "problema principal, impacto identificado, urgência e objetivo do projeto.]"
)

add_placeholder(
    "[Demonstrar compreensão do negócio do cliente sem inventar "
    "informações ou apresentar inferências como fatos.]"
)

doc.add_page_break()


# ============================================================
# 02 — SOLUÇÃO PROPOSTA
# ============================================================

add_section_header("02", "Solução proposta")

add_placeholder(
    "[Apresentar os serviços recomendados conforme o catálogo oficial "
    "da NEXORA CONSULTING.]"
)

add_placeholder(
    "[Para cada serviço, estabelecer a relação: "
    "NECESSIDADE → SERVIÇO → RESULTADO ESPERADO.]"
)

doc.add_paragraph()

p = doc.add_paragraph()
r = p.add_run("SERVIÇO 01")
r.font.name = FONT
r.font.size = Pt(9)
r.font.bold = True
r.font.color.rgb = hex_to_rgb(SECONDARY)

p = doc.add_paragraph()
r = p.add_run("[NOME OFICIAL DO SERVIÇO]")
r.font.name = FONT
r.font.size = Pt(13)
r.font.bold = True
r.font.color.rgb = hex_to_rgb(PRIMARY)

add_placeholder("[Justificativa da recomendação para este cliente.]")

doc.add_page_break()


# ============================================================
# 03 — ESCOPO E ENTREGÁVEIS
# ============================================================

add_section_header("03", "Escopo e entregáveis")

add_placeholder(
    "[Detalhar o que está incluído no projeto, utilizando somente "
    "informações validadas e entregáveis compatíveis com o catálogo.]"
)

table = doc.add_table(rows=4, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = ["ITEM", "ENTREGÁVEL", "DESCRIÇÃO"]

for col, text in enumerate(headers):
    cell = table.cell(0, col)
    set_cell_background(cell, PRIMARY)
    set_cell_border(cell, BORDER)

    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = hex_to_rgb(WHITE)

set_repeat_table_header(table.rows[0])

for row in range(1, 4):
    values = [
        f"{row:02d}",
        "[ENTREGÁVEL]",
        "[DESCRIÇÃO]"
    ]

    for col, value in enumerate(values):
        cell = table.cell(row, col)
        set_cell_border(cell, BORDER)

        if row % 2 == 0:
            set_cell_background(cell, LIGHT_BG)

        p = cell.paragraphs[0]
        r = p.add_run(value)
        r.font.name = FONT
        r.font.size = Pt(9)
        r.font.color.rgb = hex_to_rgb(TEXT)

doc.add_page_break()


# ============================================================
# 04 — CRONOGRAMA
# ============================================================

add_section_header("04", "Cronograma")

add_placeholder(
    "[Apresentar cronograma coerente com o catálogo e com a complexidade "
    "do projeto. Não somar automaticamente prazos de múltiplos serviços.]"
)

table = doc.add_table(rows=5, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = ["FASE", "PERÍODO", "PRINCIPAIS ATIVIDADES"]

for col, text in enumerate(headers):
    cell = table.cell(0, col)
    set_cell_background(cell, PRIMARY)
    set_cell_border(cell, BORDER)

    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = hex_to_rgb(WHITE)

set_repeat_table_header(table.rows[0])

phases = [
    ("01", "[PERÍODO]", "[ATIVIDADES]"),
    ("02", "[PERÍODO]", "[ATIVIDADES]"),
    ("03", "[PERÍODO]", "[ATIVIDADES]"),
    ("04", "[PERÍODO]", "[ATIVIDADES]"),
]

for row, values in enumerate(phases, start=1):
    for col, value in enumerate(values):
        cell = table.cell(row, col)
        set_cell_border(cell, BORDER)

        if row % 2 == 0:
            set_cell_background(cell, LIGHT_BG)

        p = cell.paragraphs[0]
        r = p.add_run(value)
        r.font.name = FONT
        r.font.size = Pt(9)
        r.font.color.rgb = hex_to_rgb(TEXT)

doc.add_page_break()


# ============================================================
# 05 — INVESTIMENTO
# ============================================================

add_section_header("05", "Investimento")

table = doc.add_table(rows=5, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = ["SERVIÇO", "PRAZO", "INVESTIMENTO"]

for col, text in enumerate(headers):
    cell = table.cell(0, col)
    set_cell_background(cell, PRIMARY)
    set_cell_border(cell, BORDER)

    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = hex_to_rgb(WHITE)

set_repeat_table_header(table.rows[0])

investment_rows = [
    ("[SERVIÇO 01]", "[PRAZO]", "R$ [VALOR]"),
    ("[SERVIÇO 02]", "[PRAZO]", "R$ [VALOR]"),
    ("DESCONTO", "-", "R$ [VALOR]"),
    ("INVESTIMENTO TOTAL", "[PRAZO TOTAL]", "R$ [VALOR FINAL]"),
]

for row, values in enumerate(investment_rows, start=1):
    for col, value in enumerate(values):
        cell = table.cell(row, col)
        set_cell_border(cell, BORDER)

        is_total = row == 4

        if is_total:
            set_cell_background(cell, PRIMARY)
        elif row % 2 == 0:
            set_cell_background(cell, LIGHT_BG)

        p = cell.paragraphs[0]
        r = p.add_run(value)
        r.font.name = FONT
        r.font.size = Pt(9)

        if is_total:
            r.font.bold = True
            r.font.color.rgb = hex_to_rgb(WHITE)
        else:
            r.font.color.rgb = hex_to_rgb(TEXT)

doc.add_paragraph()

p = doc.add_paragraph()
r = p.add_run("CONDIÇÕES DE PAGAMENTO")
r.font.name = FONT
r.font.size = Pt(10)
r.font.bold = True
r.font.color.rgb = hex_to_rgb(PRIMARY)

for text in [
    "30% na contratação — R$ [VALOR]",
    "40% no marco intermediário — R$ [VALOR]",
    "30% na entrega — R$ [VALOR]",
]:
    p = doc.add_paragraph(style=None)
    p.style = doc.styles["Normal"]
    r = p.add_run(f"• {text}")
    r.font.name = FONT
    r.font.size = Pt(10)
    r.font.color.rgb = hex_to_rgb(TEXT)

doc.add_page_break()


# ============================================================
# 06 — PRÓXIMOS PASSOS
# ============================================================

add_section_header("06", "Próximos passos")

next_steps = [
    "01  Revisão e aprovação da proposta",
    "02  Alinhamento final de escopo",
    "03  Formalização contratual",
    "04  Kickoff do projeto",
]

for step in next_steps:
    p = doc.add_paragraph()

    r = p.add_run(step)
    r.font.name = FONT
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = hex_to_rgb(TEXT)

    p.paragraph_format.space_after = Pt(10)

doc.add_paragraph()

p = doc.add_paragraph()

r = p.add_run("Vamos transformar esta oportunidade em execução →")
r.font.name = FONT
r.font.size = Pt(12)
r.font.bold = True
r.font.color.rgb = hex_to_rgb(SECONDARY)

doc.add_paragraph()
doc.add_paragraph()

p = doc.add_paragraph()

r = p.add_run("[NOME DO CONSULTOR]")
r.font.name = FONT
r.font.size = Pt(11)
r.font.bold = True
r.font.color.rgb = hex_to_rgb(PRIMARY)

p = doc.add_paragraph()

r = p.add_run("[CARGO]")
r.font.name = FONT
r.font.size = Pt(9)
r.font.color.rgb = hex_to_rgb(MUTED)

p = doc.add_paragraph()

r = p.add_run("NEXORA CONSULTING")
r.font.name = FONT
r.font.size = Pt(9)
r.font.bold = True
r.font.color.rgb = hex_to_rgb(TEXT)

p = doc.add_paragraph()

r = p.add_run("[EMAIL] · [TELEFONE]")
r.font.name = FONT
r.font.size = Pt(9)
r.font.color.rgb = hex_to_rgb(MUTED)


# ============================================================
# SAVE
# ============================================================

DOCS_DIR.mkdir(parents=True, exist_ok=True)

doc.save(OUTPUT_PATH)

print("=" * 60)
print("DEALCRAFT AI — TEMPLATE GERADO")
print("=" * 60)
print(f"Arquivo: {OUTPUT_PATH}")
print("Status: OK")
print("=" * 60)
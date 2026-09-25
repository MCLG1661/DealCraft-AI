from pathlib import Path
import re
import sys
import unicodedata
from datetime import datetime

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from dealcraft_validator import (
    extract_opportunity,
    load_catalog,
    calculate_readiness,
    readiness_classification,
    find_missing_mandatory,
    validate_commercial_rules,
    determine_status,
    format_brl,
    is_available,
)

from dealcraft_summary import (
    build_executive_summary,
    print_executive_summary,
    request_human_approval,
)


# ============================================================
# DEALCRAFT AI — PROPOSAL GENERATOR
# Version: 1.2
# ============================================================

BASE_DIR = Path(__file__).resolve().parent
DOCS_DIR = BASE_DIR / "DOCS"
OUTPUTS_DIR = BASE_DIR / "outputs"

BRAND_PATH = DOCS_DIR / "brand-book.json"


# ============================================================
# VISUAL CONSTANTS
# ============================================================

NAVY = "14213D"
BLUE = "2563EB"
CYAN = "06B6D4"
GRAPHITE = "1F2937"
SLATE = "64748B"
CLOUD = "F8FAFC"
BORDER = "E2E8F0"
WHITE = "FFFFFF"


# ============================================================
# TEXT UTILITIES
# ============================================================

def safe_value(value, default="NÃO INFORMADO"):
    if not is_available(value):
        return default

    return str(value).strip()


def clean_multiline(value):
    if not is_available(value):
        return ["NÃO INFORMADO"]

    result = []

    for line in str(value).splitlines():
        line = line.strip()

        if not line:
            continue

        line = line.lstrip("-–—• ").strip()

        if line:
            result.append(line)

    if not result:
        return ["NÃO INFORMADO"]

    return result


def remove_accents(text):
    normalized = unicodedata.normalize(
        "NFKD",
        str(text)
    )

    return "".join(
        char
        for char in normalized
        if not unicodedata.combining(char)
    )


def slugify_filename(text):
    text = remove_accents(text)

    text = re.sub(
        r"[^A-Za-z0-9]+",
        "-",
        text
    )

    text = text.strip("-")

    return text.upper()


# ============================================================
# WORD UTILITIES
# ============================================================

def set_cell_background(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()

    shading = OxmlElement("w:shd")

    shading.set(
        qn("w:fill"),
        color
    )

    tc_pr.append(shading)


def set_cell_border(
    cell,
    color=BORDER,
    size="6"
):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()

    tc_borders = tc_pr.first_child_found_in(
        "w:tcBorders"
    )

    if tc_borders is None:
        tc_borders = OxmlElement(
            "w:tcBorders"
        )

        tc_pr.append(tc_borders)

    for edge in (
        "top",
        "left",
        "bottom",
        "right"
    ):
        tag = f"w:{edge}"

        element = tc_borders.find(
            qn(tag)
        )

        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)

        element.set(
            qn("w:val"),
            "single"
        )

        element.set(
            qn("w:sz"),
            size
        )

        element.set(
            qn("w:color"),
            color
        )


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()

    tbl_header = OxmlElement(
        "w:tblHeader"
    )

    tbl_header.set(
        qn("w:val"),
        "true"
    )

    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    """
    Impede que uma linha da tabela seja dividida
    entre duas páginas.
    """
    tr_pr = row._tr.get_or_add_trPr()

    cant_split = OxmlElement(
        "w:cantSplit"
    )

    cant_split.set(
        qn("w:val"),
        "true"
    )

    tr_pr.append(cant_split)


def keep_with_next(paragraph):
    """
    Mantém o parágrafo junto ao próximo elemento.

    É utilizado principalmente em títulos e subtítulos
    para evitar que apareçam isolados no final de uma página.
    """
    paragraph.paragraph_format.keep_with_next = True


def keep_together(paragraph):
    """
    Evita, quando possível, que o próprio parágrafo
    seja dividido entre duas páginas.
    """
    paragraph.paragraph_format.keep_together = True


def add_page_number(paragraph):
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.RIGHT
    )

    run = paragraph.add_run(
        "Página "
    )

    run.font.size = Pt(8)

    run.font.color.rgb = (
        RGBColor.from_string(SLATE)
    )

    fld_char_1 = OxmlElement(
        "w:fldChar"
    )

    fld_char_1.set(
        qn("w:fldCharType"),
        "begin"
    )

    instr_text = OxmlElement(
        "w:instrText"
    )

    instr_text.set(
        qn("xml:space"),
        "preserve"
    )

    instr_text.text = "PAGE"

    fld_char_2 = OxmlElement(
        "w:fldChar"
    )

    fld_char_2.set(
        qn("w:fldCharType"),
        "end"
    )

    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


# ============================================================
# DOCUMENT CONFIGURATION
# ============================================================

def configure_document(document):
    section = document.sections[0]

    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)

    styles = document.styles

    normal = styles["Normal"]

    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)

    normal.font.color.rgb = (
        RGBColor.from_string(GRAPHITE)
    )

    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15

    for style_name in (
        "Title",
        "Heading 1",
        "Heading 2"
    ):
        style = styles[style_name]
        style.font.name = "Aptos"

    title = styles["Title"]

    title.font.size = Pt(30)
    title.font.bold = True

    title.font.color.rgb = (
        RGBColor.from_string(NAVY)
    )

    heading_1 = styles["Heading 1"]

    heading_1.font.size = Pt(17)
    heading_1.font.bold = True

    heading_1.font.color.rgb = (
        RGBColor.from_string(NAVY)
    )

    heading_1.paragraph_format.keep_with_next = True
    heading_1.paragraph_format.keep_together = True

    heading_2 = styles["Heading 2"]

    heading_2.font.size = Pt(12)
    heading_2.font.bold = True

    heading_2.font.color.rgb = (
        RGBColor.from_string(BLUE)
    )

    heading_2.paragraph_format.keep_with_next = True
    heading_2.paragraph_format.keep_together = True

    header = section.header

    paragraph = header.paragraphs[0]

    paragraph.text = (
        "NEXORA CONSULTING  |  "
        "Data · AI · Growth · Automation"
    )

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.RIGHT
    )

    for run in paragraph.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.bold = True

        run.font.color.rgb = (
            RGBColor.from_string(SLATE)
        )

    footer = section.footer

    paragraph = footer.paragraphs[0]

    paragraph.text = (
        "NEXORA CONSULTING  •  "
        "Proposta Comercial Confidencial"
    )

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT
    )

    for run in paragraph.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)

        run.font.color.rgb = (
            RGBColor.from_string(SLATE)
        )

    page_paragraph = footer.add_paragraph()

    add_page_number(
        page_paragraph
    )


# ============================================================
# DOCUMENT COMPONENTS
# ============================================================

def add_section_title(
    document,
    number,
    title
):
    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_before = (
        Pt(18)
    )

    paragraph.paragraph_format.space_after = (
        Pt(8)
    )

    keep_with_next(paragraph)
    keep_together(paragraph)

    run_number = paragraph.add_run(
        f"{number}  "
    )

    run_number.font.name = "Aptos"
    run_number.font.size = Pt(17)
    run_number.font.bold = True

    run_number.font.color.rgb = (
        RGBColor.from_string(CYAN)
    )

    run_title = paragraph.add_run(
        title
    )

    run_title.font.name = "Aptos"
    run_title.font.size = Pt(17)
    run_title.font.bold = True

    run_title.font.color.rgb = (
        RGBColor.from_string(NAVY)
    )

    return paragraph


def add_subheading(
    document,
    title
):
    paragraph = document.add_heading(
        title,
        level=2
    )

    keep_with_next(paragraph)
    keep_together(paragraph)

    return paragraph


def add_label_value(
    document,
    label,
    value
):
    paragraph = document.add_paragraph()

    keep_together(paragraph)

    label_run = paragraph.add_run(
        f"{label}: "
    )

    label_run.bold = True

    label_run.font.color.rgb = (
        RGBColor.from_string(GRAPHITE)
    )

    value_run = paragraph.add_run(
        safe_value(value)
    )

    value_run.font.color.rgb = (
        RGBColor.from_string(GRAPHITE)
    )

    return paragraph


def add_bullets(
    document,
    values
):
    paragraphs = []

    for value in clean_multiline(values):
        paragraph = document.add_paragraph(
            style="List Bullet"
        )

        paragraph.add_run(
            value
        )

        keep_together(paragraph)

        paragraphs.append(paragraph)

    return paragraphs


def add_info_table(
    document,
    rows
):
    table = document.add_table(
        rows=0,
        cols=2
    )

    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    table.autofit = True

    for label, value in rows:
        row = table.add_row()

        set_row_cant_split(row)

        cells = row.cells

        cells[0].text = label
        cells[1].text = safe_value(value)

        set_cell_background(
            cells[0],
            CLOUD
        )

        for cell in cells:
            set_cell_border(cell)

            for paragraph in cell.paragraphs:
                keep_together(paragraph)

                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(9.5)

        for run in cells[0].paragraphs[0].runs:
            run.bold = True

            run.font.color.rgb = (
                RGBColor.from_string(NAVY)
            )

    return table


# ============================================================
# COVER
# ============================================================

def add_cover(
    document,
    opportunity,
    commercial
):
    document.add_paragraph()

    brand = document.add_paragraph()

    brand.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT
    )

    run = brand.add_run(
        "NEXORA"
    )

    run.font.name = "Aptos"
    run.font.size = Pt(14)
    run.font.bold = True

    run.font.color.rgb = (
        RGBColor.from_string(BLUE)
    )

    subtitle = document.add_paragraph()

    run = subtitle.add_run(
        "CONSULTING"
    )

    run.font.name = "Aptos"
    run.font.size = Pt(9)
    run.font.bold = True

    run.font.color.rgb = (
        RGBColor.from_string(SLATE)
    )

    document.add_paragraph()
    document.add_paragraph()
    document.add_paragraph()

    title = document.add_paragraph(
        style="Title"
    )

    keep_with_next(title)
    keep_together(title)

    title.add_run(
        "Proposta Comercial"
    )

    client = document.add_paragraph()

    keep_with_next(client)
    keep_together(client)

    client_run = client.add_run(
        safe_value(
            opportunity.get("A1")
        )
    )

    client_run.font.name = "Aptos"
    client_run.font.size = Pt(20)
    client_run.font.bold = True

    client_run.font.color.rgb = (
        RGBColor.from_string(BLUE)
    )

    document.add_paragraph()

    service = document.add_paragraph()

    keep_together(service)

    service_run = service.add_run(
        safe_value(
            opportunity.get("D1")
        )
    )

    service_run.font.name = "Aptos"
    service_run.font.size = Pt(14)

    service_run.font.color.rgb = (
        RGBColor.from_string(GRAPHITE)
    )

    document.add_paragraph()
    document.add_paragraph()

    proposal_number = safe_value(
        opportunity.get("E1")
    )

    issue_date = safe_value(
        opportunity.get("E4")
    )

    consultant = safe_value(
        opportunity.get("E2")
    )

    add_info_table(
        document,
        [
            (
                "Proposta",
                proposal_number
            ),
            (
                "Data de emissão",
                issue_date
            ),
            (
                "Consultor responsável",
                consultant
            ),
            (
                "Investimento",
                format_brl(
                    commercial["final_value"]
                )
            ),
            (
                "Validade",
                "30 dias corridos"
            ),
        ]
    )

    document.add_paragraph()

    note = document.add_paragraph()

    keep_together(note)

    note.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT
    )

    run = note.add_run(
        "Transforme oportunidades comerciais "
        "em propostas estruturadas, consistentes "
        "e prontas para decisão."
    )

    run.font.name = "Aptos"
    run.font.size = Pt(10)
    run.font.italic = True

    run.font.color.rgb = (
        RGBColor.from_string(SLATE)
    )

    document.add_page_break()


# ============================================================
# SECTION 01
# ============================================================

def add_context_section(
    document,
    opportunity
):
    add_section_title(
        document,
        "01",
        "Contexto e oportunidade"
    )

    add_subheading(
        document,
        "Cenário atual"
    )

    for paragraph_text in clean_multiline(
        opportunity.get("B1")
    ):
        paragraph = document.add_paragraph(
            paragraph_text
        )

        keep_together(paragraph)

    add_subheading(
        document,
        "Principal problema"
    )

    for paragraph_text in clean_multiline(
        opportunity.get("B2")
    ):
        paragraph = document.add_paragraph(
            paragraph_text
        )

        keep_together(paragraph)

    add_subheading(
        document,
        "Objetivo do projeto"
    )

    for paragraph_text in clean_multiline(
        opportunity.get("B3")
    ):
        paragraph = document.add_paragraph(
            paragraph_text
        )

        keep_together(paragraph)

    add_subheading(
        document,
        "Impactos identificados"
    )

    add_bullets(
        document,
        opportunity.get("B4")
    )


# ============================================================
# SECTION 02
# ============================================================

def add_solution_section(
    document,
    opportunity,
    commercial
):
    add_section_title(
        document,
        "02",
        "Solução proposta"
    )

    service = commercial["service"]

    add_subheading(
        document,
        safe_value(
            opportunity.get("D1")
        )
    )

    if service:
        paragraph = document.add_paragraph(
            service.get(
                "description",
                "NÃO INFORMADO"
            )
        )

        keep_together(paragraph)

    add_subheading(
        document,
        "Objetivo da solução"
    )

    for paragraph_text in clean_multiline(
        opportunity.get("B3")
    ):
        paragraph = document.add_paragraph(
            paragraph_text
        )

        keep_together(paragraph)

    if service:
        deliverables = service.get(
            "deliverables",
            []
        )

        if deliverables:
            add_subheading(
                document,
                "Entregáveis de referência"
            )

            for deliverable in deliverables:
                paragraph = (
                    document.add_paragraph(
                        style="List Bullet"
                    )
                )

                paragraph.add_run(
                    deliverable
                )

                keep_together(paragraph)


# ============================================================
# SECTION 03
# ============================================================

def add_scope_section(
    document,
    opportunity
):
    add_section_title(
        document,
        "03",
        "Escopo e entregáveis"
    )

    add_subheading(
        document,
        "Escopo acordado"
    )

    add_bullets(
        document,
        opportunity.get("D2")
    )

    add_subheading(
        document,
        "Premissas"
    )

    add_bullets(
        document,
        opportunity.get("D6")
    )

    add_subheading(
        document,
        "Exclusões"
    )

    add_bullets(
        document,
        opportunity.get("D7")
    )


# ============================================================
# SECTION 04
# ============================================================

def add_timeline_section(
    document,
    opportunity,
    commercial
):
    add_section_title(
        document,
        "04",
        "Cronograma"
    )

    service = commercial["service"]

    add_info_table(
        document,
        [
            (
                "Prazo proposto",
                opportunity.get("D3")
            ),
            (
                "Prazo desejado pelo cliente",
                opportunity.get("B6")
            ),
            (
                "Urgência",
                opportunity.get("B5")
            ),
        ]
    )

    if service:
        duration = service[
            "duration"
        ]

        document.add_paragraph()

        paragraph = document.add_paragraph()

        keep_together(paragraph)

        run = paragraph.add_run(
            "Faixa de referência do catálogo: "
        )

        run.bold = True

        paragraph.add_run(
            f"{duration['minimum_days']} a "
            f"{duration['maximum_days']} dias."
        )


# ============================================================
# SECTION 05
# ============================================================

def add_investment_section(
    document,
    commercial
):
    add_section_title(
        document,
        "05",
        "Investimento"
    )

    rows = [
        (
            "Investimento bruto",
            format_brl(
                commercial["gross_value"]
            )
        ),
        (
            "Desconto",
            f"{commercial['discount_percent']}%"
        ),
        (
            "Valor do desconto",
            format_brl(
                commercial["discount_value"]
            )
        ),
        (
            "Investimento final",
            format_brl(
                commercial["final_value"]
            )
        ),
    ]

    add_info_table(
        document,
        rows
    )

    document.add_paragraph()

    add_subheading(
        document,
        "Condições de pagamento"
    )

    table = document.add_table(
        rows=1,
        cols=3
    )

    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    headers = (
        "Parcela",
        "Evento",
        "Valor"
    )

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]

        cell.text = header

        set_cell_background(
            cell,
            NAVY
        )

        set_cell_border(
            cell
        )

        for paragraph in cell.paragraphs:
            keep_together(paragraph)

        for run in cell.paragraphs[0].runs:
            run.bold = True

            run.font.color.rgb = (
                RGBColor.from_string(WHITE)
            )

    # A tabela possui poucas linhas.
    # Não repetimos automaticamente o cabeçalho para evitar
    # um cabeçalho isolado caso a tabela alcance o limite
    # inferior da página.

    set_row_cant_split(
        table.rows[0]
    )

    for installment in commercial[
        "installments"
    ]:
        row = table.add_row()

        set_row_cant_split(
            row
        )

        cells = row.cells

        cells[0].text = (
            f"{installment['percentage']}%"
        )

        cells[1].text = (
            installment["event"].capitalize()
        )

        cells[2].text = format_brl(
            installment["amount"]
        )

        for cell in cells:
            set_cell_border(cell)

            for paragraph in cell.paragraphs:
                keep_together(paragraph)

                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(9.5)

    document.add_paragraph()

    paragraph = document.add_paragraph()

    keep_together(paragraph)

    run = paragraph.add_run(
        "Validade da proposta: "
    )

    run.bold = True

    paragraph.add_run(
        "30 dias corridos."
    )


# ============================================================
# SECTION 06
# ============================================================

def add_next_steps_section(
    document,
    opportunity
):
    add_section_title(
        document,
        "06",
        "Próximos passos"
    )

    add_label_value(
        document,
        "Próximo passo comercial",
        opportunity.get("C5")
    )

    add_label_value(
        document,
        "Data prevista",
        opportunity.get("C6")
    )

    add_label_value(
        document,
        "Decisor",
        opportunity.get("C1")
    )

    document.add_paragraph()

    paragraph = document.add_paragraph()

    keep_together(paragraph)

    run = paragraph.add_run(
        "Após a aprovação comercial, "
        "as partes poderão avançar para "
        "formalização e planejamento do início "
        "do projeto, respeitando as condições "
        "registradas nesta proposta."
    )

    run.font.color.rgb = (
        RGBColor.from_string(GRAPHITE)
    )

    document.add_paragraph()
    document.add_paragraph()

    closing = document.add_paragraph()

    keep_with_next(closing)
    keep_together(closing)

    closing.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = closing.add_run(
        "NEXORA CONSULTING"
    )

    run.bold = True
    run.font.size = Pt(12)

    run.font.color.rgb = (
        RGBColor.from_string(NAVY)
    )

    closing_2 = document.add_paragraph()

    keep_together(closing_2)

    closing_2.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = closing_2.add_run(
        "Data · AI · Growth · Automation"
    )

    run.font.size = Pt(9)

    run.font.color.rgb = (
        RGBColor.from_string(SLATE)
    )


# ============================================================
# FILENAME
# ============================================================

def build_filename(opportunity):
    proposal_number = slugify_filename(
        safe_value(
            opportunity.get("E1"),
            "SEM-NUMERO"
        )
    )

    client = slugify_filename(
        safe_value(
            opportunity.get("A1"),
            "CLIENTE"
        )
    )

    issue_date = safe_value(
        opportunity.get("E4")
    )

    year_match = re.search(
        r"\b(20\d{2})\b",
        issue_date
    )

    if year_match:
        year = year_match.group(1)

    else:
        year = str(
            datetime.now().year
        )

    filename = (
        f"Proposta-"
        f"{proposal_number}-"
        f"{year}-"
        f"{client}-"
        f"NEXORA-CONSULTING.docx"
    )

    return filename


# ============================================================
# DOCUMENT GENERATOR
# ============================================================

def generate_proposal(
    opportunity,
    commercial
):
    OUTPUTS_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    document = Document()

    configure_document(
        document
    )

    add_cover(
        document,
        opportunity,
        commercial
    )

    add_context_section(
        document,
        opportunity
    )

    add_solution_section(
        document,
        opportunity,
        commercial
    )

    add_scope_section(
        document,
        opportunity
    )

    add_timeline_section(
        document,
        opportunity,
        commercial
    )

    add_investment_section(
        document,
        commercial
    )

    add_next_steps_section(
        document,
        opportunity
    )

    filename = build_filename(
        opportunity
    )

    output_path = (
        OUTPUTS_DIR / filename
    )

    document.save(
        output_path
    )

    return output_path


# ============================================================
# PROCESS
# ============================================================

def process_proposal(input_path):
    if not input_path.is_absolute():
        input_path = (
            BASE_DIR / input_path
        )

    input_path = input_path.resolve()

    if not input_path.exists():
        print()

        print(
            f"ERRO: arquivo não encontrado: "
            f"{input_path}"
        )

        return 1

    try:
        text = input_path.read_text(
            encoding="utf-8"
        )

        catalog = load_catalog()

    except Exception as error:
        print()

        print(
            f"ERRO: {error}"
        )

        return 1

    opportunity = extract_opportunity(
        text
    )

    (
        dimensions,
        available,
        total,
        percentage
    ) = calculate_readiness(
        opportunity
    )

    readiness_label = (
        readiness_classification(
            percentage
        )
    )

    missing_mandatory = (
        find_missing_mandatory(
            opportunity
        )
    )

    commercial = (
        validate_commercial_rules(
            opportunity,
            catalog
        )
    )

    status = determine_status(
        missing_mandatory,
        commercial
    )

    # --------------------------------------------------------
    # VALIDATION GATE
    # --------------------------------------------------------

    if status != (
        "APROVADO PARA RESUMO EXECUTIVO"
    ):
        print()
        print("=" * 78)

        print(
            "DEALCRAFT AI — PROPOSAL GENERATOR"
        )

        print("=" * 78)

        print()

        print(
            f"Status da oportunidade: "
            f"{status}"
        )

        print()

        print(
            "GERAÇÃO DA PROPOSTA: "
            "NÃO AUTORIZADA"
        )

        if missing_mandatory:
            print()

            print(
                "Campos obrigatórios pendentes:"
            )

            for item in missing_mandatory:
                print(
                    f"- {item['field']} — "
                    f"{item['description']}"
                )

        print()
        print("=" * 78)

        return 2

    # --------------------------------------------------------
    # EXECUTIVE SUMMARY
    # --------------------------------------------------------

    summary = build_executive_summary(
        opportunity,
        commercial
    )

    print_executive_summary(
        summary,
        percentage,
        readiness_label
    )

    # --------------------------------------------------------
    # HUMAN APPROVAL
    # --------------------------------------------------------

    approved = request_human_approval()

    if not approved:
        print()
        print("=" * 78)

        print(
            "APROVAÇÃO NÃO CONCEDIDA"
        )

        print("=" * 78)

        print()

        print(
            "GERAÇÃO DA PROPOSTA: "
            "CANCELADA"
        )

        print()

        print(
            "Nenhum arquivo foi criado."
        )

        print("=" * 78)

        return 3

    # --------------------------------------------------------
    # GENERATION
    # --------------------------------------------------------

    print()

    print(
        "Aprovação humana registrada."
    )

    print()

    print(
        "Gerando proposta comercial..."
    )

    try:
        output_path = generate_proposal(
            opportunity,
            commercial
        )

    except Exception as error:
        print()

        print(
            "ERRO DURANTE A GERAÇÃO "
            "DA PROPOSTA:"
        )

        print(error)

        return 4

    print()
    print("=" * 78)

    print(
        "PROPOSTA GERADA COM SUCESSO"
    )

    print("=" * 78)

    print()

    print(
        f"Arquivo: {output_path.name}"
    )

    print()

    print(
        f"Local: {output_path}"
    )

    print()

    print(
        "Status: DOCUMENTO GERADO"
    )

    print("=" * 78)

    return 0


# ============================================================
# CLI
# ============================================================

def print_usage():
    print()

    print(
        "DEALCRAFT AI — PROPOSAL GENERATOR 1.2"
    )

    print()
    print("Uso:")

    print(
        "python dealcraft_proposal.py "
        ".\\examples\\briefing-completo.txt"
    )


def main():
    if len(sys.argv) < 2:
        print_usage()
        return 1

    input_path = Path(
        sys.argv[1]
    )

    return process_proposal(
        input_path
    )


if __name__ == "__main__":
    sys.exit(
        main()
    )